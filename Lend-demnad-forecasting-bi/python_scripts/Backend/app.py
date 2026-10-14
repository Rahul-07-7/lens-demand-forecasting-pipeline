from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from fastapi import Body
from forecast_service import run_forecast
import pandas as pd
import io
import traceback
import logging
from io import BytesIO
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    Alignment,
    PatternFill,
    Border,
    Side
)

app = FastAPI()

logging.basicConfig(level=logging.INFO)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"], 
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

uploaded_data = None
last_forecast = None


def read_uploaded_file(contents, filename):

    if filename.endswith(".csv"):

        try:
            return pd.read_csv(io.BytesIO(contents))

        except Exception:
            return pd.read_csv(
                io.BytesIO(contents),
                encoding="latin1"
            )

    return pd.read_excel(io.BytesIO(contents))


def clean_columns(df):
    df.columns = [c.lower().strip() for c in df.columns]
    return df


def validate_columns(df, required_cols):

    missing = [c for c in required_cols if c not in df.columns]

    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Missing columns: {missing}"
        )


class AllocateRequest(BaseModel):
    forecast: list
    sub_products: list


class AllocateSubSubRequest(BaseModel):
    sub_alloc: list
    subsub_data: list


class HierarchicalRequest(BaseModel):
    forecast: list
    mapping_data: list


@app.get("/")
def home():
    return {
        "status": "running",
        "message": "Forecast API is working"
    }


@app.post("/upload-and-forecast")
async def upload_and_forecast(
    file: UploadFile = File(...)
):

    global uploaded_data, last_forecast

    try:

        contents = await file.read()

        df = read_uploaded_file(
            contents,
            file.filename
        )

        df = clean_columns(df)

        validate_columns(
            df,
            ["date", "sales"]
        )

        df = df.dropna(subset=["date", "sales"])

        uploaded_data = df[
            ["date", "sales"]
        ].to_dict(orient="records")

        result = run_forecast(uploaded_data)

        last_forecast = result

        return result

    except Exception as e:

        traceback.print_exc()

        raise HTTPException(
            status_code=400,
            detail=str(e)
        )


@app.post("/calculate-std")
async def calculate_std(
    file: UploadFile = File(...)
):

    try:

        contents = await file.read()

        df = read_uploaded_file(
            contents,
            file.filename
        )

        df = clean_columns(df)

        if "quantity" in df.columns:
            df.rename(
                columns={"quantity": "sales"},
                inplace=True
            )

        validate_columns(
            df,
            ["product", "sales"]
        )

        df["sales"] = pd.to_numeric(
            df["sales"],
            errors="coerce"
        )

        df = df.dropna(
            subset=["product", "sales"]
        )

        result = []

        for product, group in df.groupby("product"):

            values = group[
                group["sales"] > 0
            ]["sales"]

            if len(values) > 1:
                std_dev = float(
                    values.std(ddof=1)
                )
            else:
                std_dev = 0.0

            result.append({
                "product": str(product).strip(),
                "std_dev": round(std_dev, 2)
            })

        return result

    except Exception as e:

        traceback.print_exc()

        raise HTTPException(
            status_code=400,
            detail=str(e)
        )


@app.post("/calculate-std-pivot")
async def calculate_std_pivot(
    file: UploadFile = File(...)
):

    try:

        contents = await file.read()

        df = read_uploaded_file(
            contents,
            file.filename
        )

        df = clean_columns(df)

        if "quantity" in df.columns:
            df.rename(
                columns={"quantity": "sales"},
                inplace=True
            )

        validate_columns(
            df,
            ["power", "powsph", "sales"]
        )

        df["sales"] = pd.to_numeric(
            df["sales"],
            errors="coerce"
        )

        df = df.dropna(
            subset=["power", "powsph", "sales"]
        )

        result = []

        for (power, powsph), group in df.groupby(
            ["power", "powsph"]
        ):

            values = group[
                group["sales"] > 0
            ]["sales"]

            if len(values) > 1:
                std_dev = float(
                    values.std(ddof=1)
                )
            else:
                std_dev = 0.0

            if pd.isna(std_dev):
                std_dev = 0.0

            result.append({
                "Power": str(power),
                "PowSph": str(powsph),
                "std_dev": round(std_dev, 2)
            })

        return result

    except Exception as e:

        traceback.print_exc()

        raise HTTPException(
            status_code=400,
            detail=str(e)
        )


@app.post("/upload-current-stock")
async def upload_stock(
    file: UploadFile = File(...)
):

    try:

        contents = await file.read()

        df = read_uploaded_file(
            contents,
            file.filename
        )

        df = clean_columns(df)

        validate_columns(
            df,
            ["product","current_stock"]
        )

        df["current_stock"] = pd.to_numeric(
            df["current_stock"],
            errors="coerce"
        ).fillna(0)

        result = []

        for _, row in df.iterrows():

            result.append({
                "product": str(
                    row["product"]
                ).strip(),

                "current_stock": float(
                    row["current_stock"]
                )
            })

        return result

    except Exception as e:

        traceback.print_exc()

        raise HTTPException(
            status_code=400,
            detail=str(e)
        )


def prepare_subproduct_data(df):

    df = clean_columns(df)

    validate_columns(
        df,
        ["product", "total_sales"]
    )

    df["total_sales"] = pd.to_numeric(
        df["total_sales"],
        errors="coerce"
    ).fillna(0)

    return df.to_dict(orient="records")


@app.post("/upload-subproducts")
async def upload_subproducts(
    file: UploadFile = File(...)
):

    contents = await file.read()

    df = read_uploaded_file(
        contents,
        file.filename
    )

    return prepare_subproduct_data(df)


@app.post("/upload-subsub")
async def upload_subsub(
    file: UploadFile = File(...)
):

    contents = await file.read()

    df = read_uploaded_file(
        contents,
        file.filename
    )

    df.columns = [
        c.strip().upper()
        for c in df.columns
    ]

    required = {"SUBP", "SUBP2", "SALES"}

    if not required.issubset(df.columns):

        raise HTTPException(
            status_code=400,
            detail="File must contain: SUBP, SUBP2, SALES"
        )

    df["SALES"] = pd.to_numeric(
        df["SALES"],
        errors="coerce"
    ).fillna(0)

    return df.to_dict(orient="records")


def allocate_yearly(
    forecast,
    sub_products
):

    df = pd.DataFrame(sub_products)

    df["total_sales"] = pd.to_numeric(
        df.get("total_sales", 0),
        errors="coerce"
    ).fillna(0)

    total_hist = df["total_sales"].sum()

    if total_hist == 0:
        raise ValueError(
            "Total historical sales is zero"
        )

    total_forecast = sum(
        f.get("forecast_sales", 0)
        for f in forecast
    )

    allocation = []

    for _, row in df.iterrows():

        share = row["total_sales"] / total_hist

        allocation.append({
            "product": str(row["product"]),
            "yearly_forecast": round(
                total_forecast * share,
                2
            )
        })

    return {
        "total_year_forecast": round(
            total_forecast,
            2
        ),
        "product_allocation": allocation
    }


@app.post("/allocate")
async def allocate(
    data: AllocateRequest
):

    result = allocate_yearly(
        data.forecast,
        data.sub_products
    )

    return result



def allocate_subsub(
    sub_alloc,
    subsub_data
):

    df = pd.DataFrame(subsub_data)

    df.columns = [
        c.strip().upper()
        for c in df.columns
    ]

    required = {"SUBP", "SUBP2", "SALES"}

    if not required.issubset(df.columns):

        raise HTTPException(
            status_code=400,
            detail=f"Missing columns: {required}"
        )

    df["SALES"] = pd.to_numeric(
        df["SALES"],
        errors="coerce"
    ).fillna(0)

    results = []

    for item in sub_alloc:

        parent_name = str(
            item.get("product", "")
        )

        parent_forecast = float(
            item.get("yearly_forecast", 0)
        )

        parent_df = df[
            df["SUBP"].astype(str)
            == parent_name
        ]

        if parent_df.empty:
            continue

        child_summary = parent_df.groupby(
            "SUBP2"
        )["SALES"].sum().reset_index()

        parent_total = (
            child_summary["SALES"].sum()
            or 1
        )

        for _, row in child_summary.iterrows():

            share = row["SALES"] / parent_total

            results.append({
                "sub_product": parent_name,
                "sub_sub_product": str(
                    row["SUBP2"]
                ),
                "forecast": round(
                    parent_forecast * share,
                    2
                )
            })

    return {
        "subsub_allocation": results
    }


@app.post("/allocate-subsub")
async def allocate_subsub_api(
    data: AllocateSubSubRequest
):

    result = allocate_subsub(
        data.sub_alloc,
        data.subsub_data
    )

    return result


@app.post("/allocate-hierarchical")
async def allocate_hierarchical(
    data: HierarchicalRequest
):

    try:

        forecast_list = data.forecast
        mapping_data = data.mapping_data

        df = pd.DataFrame(mapping_data)

        df.columns = [
            c.strip().upper()
            for c in df.columns
        ]

        validate_columns(
            clean_columns(df),
            ["subp", "subp2", "sales"]
        )

        total_f = sum(
            float(
                f.get("forecast_sales", 0)
            )
            for f in forecast_list
        )

        total_hist = float(
            df["sales"].sum()
        )

        if total_hist == 0:

            raise HTTPException(
                status_code=400,
                detail="Historical sales total is zero"
            )

        grouped = df.groupby(
            ["subp", "subp2"]
        )["sales"].sum().reset_index()

        subsub_results = []

        for _, row in grouped.iterrows():

            share = row["sales"] / total_hist

            subsub_results.append({
                "sub_product": str(
                    row["subp"]
                ),

                "sub_sub_product": str(
                    row["subp2"]
                ),

                "forecast": round(
                    total_f * share,
                    2
                )
            })

        sub_summary = pd.DataFrame(
            subsub_results
        ).groupby(
            "sub_product"
        )["forecast"].sum().reset_index()

        sub_alloc = []

        for _, row in sub_summary.iterrows():

            sub_alloc.append({
                "product": str(
                    row["sub_product"]
                ),

                "yearly_forecast": round(
                    float(row["forecast"]),
                    2
                )
            })

        return {
            "sub_allocation": sub_alloc,
            "subsub_allocation": subsub_results
        }

    except Exception as e:

        traceback.print_exc()

        raise HTTPException(
            status_code=400,
            detail=str(e)
        )


@app.get("/download-excel")
def download_excel():

    global last_forecast

    if not last_forecast or "forecast" not in last_forecast:

        raise HTTPException(
            status_code=400,
            detail="No forecast data available"
        )

    wb = Workbook()

    ws = wb.active

    ws.title = "Forecast Report"

    start_row = 2
    start_col = 2

    headers = [
        "Date",
        "Forecast Sales",
        "Min Sales",
        "Max Sales"
    ]

    # STYLES
    header_font = Font(
        bold=True,
        size=16
    )

    data_font = Font(size=14)

    header_fill = PatternFill(
        start_color="D9E1F2",
        end_color="D9E1F2",
        fill_type="solid"
    )

    total_fill = PatternFill(
        start_color="FFF2CC",
        end_color="FFF2CC",
        fill_type="solid"
    )

    center_align = Alignment(
        horizontal="center",
        vertical="center"
    )

    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # HEADERS
    for col_idx, header in enumerate(
        headers,
        start=start_col
    ):

        cell = ws.cell(
            row=start_row,
            column=col_idx,
            value=header
        )

        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border

    total_forecast = 0
    total_min = 0
    total_max = 0

    # DATA
    for row_idx, item in enumerate(
        last_forecast["forecast"],
        start=start_row + 1
    ):

        forecast_sales = item.get(
            "forecast_sales",
            0
        )

        min_sales = item.get(
            "min_sales",
            0
        )

        max_sales = item.get(
            "max_sales",
            0
        )

        total_forecast += forecast_sales
        total_min += min_sales
        total_max += max_sales

        values = [
            item.get("date"),
            forecast_sales,
            min_sales,
            max_sales
        ]

        for col_idx, value in enumerate(
            values,
            start=start_col
        ):

            cell = ws.cell(
                row=row_idx,
                column=col_idx,
                value=value
            )

            cell.font = data_font
            cell.alignment = center_align
            cell.border = border

    # TOTAL ROW
    total_row = (
        start_row
        + len(last_forecast["forecast"])
        + 1
    )

    totals = [
        "TOTAL",
        total_forecast,
        total_min,
        total_max
    ]

    for col_idx, value in enumerate(
        totals,
        start=start_col
    ):

        cell = ws.cell(
            row=total_row,
            column=col_idx,
            value=value
        )

        cell.font = Font(
            bold=True,
            size=14
        )

        cell.fill = total_fill
        cell.alignment = center_align
        cell.border = border

    # DATE FORMAT
    for row in ws.iter_rows(
        min_row=start_row + 1,
        max_row=total_row - 1,
        min_col=start_col,
        max_col=start_col
    ):

        for cell in row:

            try:

                cell.value = datetime.strptime(
                    str(cell.value),
                    "%Y-%m-%d"
                )

                cell.number_format = "MMM-YYYY"

            except Exception:
                pass

    # AUTO WIDTH
    for column_cells in ws.columns:

        max_length = max(
            (
                len(str(cell.value))
                for cell in column_cells
                if cell.value is not None
            ),
            default=0
        )

        if len(column_cells) > 0:

            col_letter = (
                column_cells[0]
                .column_letter
            )

            ws.column_dimensions[
                col_letter
            ].width = max_length + 5

    # STREAM
    stream = BytesIO()

    wb.save(stream)

    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":
            "attachment; filename=forecast.xlsx"
        }
    )


@app.post("/download-allocation-excel")
def download_allocation_excel(data: dict = Body(...)):

    from openpyxl import Workbook
    from openpyxl.styles import (
        Font,
        PatternFill,
        Alignment,
        Border,
        Side
    )

    from openpyxl.utils import get_column_letter

    from fastapi.responses import StreamingResponse

    from io import BytesIO

    try:

        allocation = data.get("allocation", [])

        selected_product = data.get(
            "selectedProduct",
            "allocation"
        )

        # =========================================
        # CREATE WORKBOOK
        # =========================================

        wb = Workbook()

        ws = wb.active

        ws.title = "Allocation Report"

        # =========================================
        # TITLE
        # =========================================

        ws.merge_cells("A1:B1")

        title_cell = ws["A1"]

        title_cell.value = (
            f"{selected_product} Allocation Report"
        )

        title_cell.font = Font(
            bold=True,
            size=20
        )

        title_cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        ws.row_dimensions[1].height = 35

        # =========================================
        # HEADERS
        # =========================================

        headers = [
            "Product",
            "Yearly Forecast"
        ]

        # =========================================
        # STYLES
        # =========================================

        header_font = Font(
            bold=True,
            size=16
        )

        data_font = Font(
            size=14
        )

        total_font = Font(
            bold=True,
            size=15
        )

        header_fill = PatternFill(
            start_color="D9E1F2",
            end_color="D9E1F2",
            fill_type="solid"
        )

        total_fill = PatternFill(
            start_color="FFF2CC",
            end_color="FFF2CC",
            fill_type="solid"
        )

        center_align = Alignment(
            horizontal="center",
            vertical="center"
        )

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        # =========================================
        # HEADER ROW
        # =========================================

        header_row = 3

        for col, header in enumerate(
            headers,
            start=1
        ):

            cell = ws.cell(
                row=header_row,
                column=col,
                value=header
            )

            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        # =========================================
        # DATA ROWS
        # =========================================

        total = 0

        start_data_row = 4

        for row_idx, item in enumerate(
            allocation,
            start=start_data_row
        ):

            product = str(
                item.get("product", "")
            )

            yearly = float(
                item.get(
                    "yearly_forecast",
                    0
                )
            )

            total += yearly

            product_cell = ws.cell(
                row=row_idx,
                column=1,
                value=product
            )

            value_cell = ws.cell(
                row=row_idx,
                column=2,
                value=round(yearly, 2)
            )

            for cell in [product_cell, value_cell]:

                cell.font = data_font
                cell.alignment = center_align
                cell.border = thin_border

        # =========================================
        # TOTAL ROW
        # =========================================

        total_row = start_data_row + len(allocation)

        total_label = ws.cell(
            row=total_row,
            column=1,
            value="TOTAL"
        )

        total_value = ws.cell(
            row=total_row,
            column=2,
            value=round(total, 2)
        )

        for cell in [total_label, total_value]:

            cell.font = total_font
            cell.fill = total_fill
            cell.alignment = center_align
            cell.border = thin_border

        # =========================================
        # AUTO COLUMN WIDTH
        # =========================================

        for col_num in range(1, ws.max_column + 1):

            column_letter = get_column_letter(col_num)

            max_length = 0

            for row in range(1, ws.max_row + 1):

                cell = ws.cell(
                    row=row,
                    column=col_num
                )

                try:

                    if cell.value:

                        max_length = max(
                            max_length,
                            len(str(cell.value))
                        )

                except:
                    pass

            adjusted_width = max_length + 10

            ws.column_dimensions[
                column_letter
            ].width = adjusted_width

        # =========================================
        # ROW HEIGHT
        # =========================================

        for row in range(1, ws.max_row + 1):

            ws.row_dimensions[row].height = 28

        # =========================================
        # SAVE FILE
        # =========================================

        stream = BytesIO()

        wb.save(stream)

        stream.seek(0)

        return StreamingResponse(
            stream,

            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

            headers={
                "Content-Disposition":
                f"attachment; filename={selected_product}_allocation.xlsx"
            }
        )

    except Exception as e:

        print("EXCEL EXPORT ERROR:", str(e))

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )

@app.post("/download-inventory-excel")
def download_inventory_excel(data: dict = Body(...)):

    from openpyxl import Workbook

    from openpyxl.styles import (
        Font,
        PatternFill,
        Alignment,
        Border,
        Side
    )

    from openpyxl.utils import (
        get_column_letter
    )

    from fastapi.responses import (
        StreamingResponse
    )

    from io import BytesIO

    try:

        inventory = data.get(
            "inventory",
            []
        )

        selected_product = data.get(
            "selectedProduct",
            "inventory"
        )

        # ====================================
        # WORKBOOK
        # ====================================

        wb = Workbook()

        ws = wb.active

        ws.title = "Inventory Planning"

        # ====================================
        # TITLE
        # ====================================

        ws.merge_cells("A1:L1")

        title = ws["A1"]

        title.value = (
            f"{selected_product} Inventory Planning"
        )

        title.font = Font(
            bold=True,
            size=22
        )

        title.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        ws.row_dimensions[1].height = 35

        # ====================================
        # HEADERS
        # ====================================

        headers = [
            "Product",
            "Forecast",
            "Buffered",
            "Band",
            "Z",
            "Std Dev",
            "Safety",
            "Min",
            "Max",
            
        ]

        # ====================================
        # STYLES
        # ====================================

        header_font = Font(
            bold=True,
            size=15
        )

        data_font = Font(
            size=13
        )

        total_font = Font(
            bold=True,
            size=14
        )

        header_fill = PatternFill(
            start_color="D9E1F2",
            end_color="D9E1F2",
            fill_type="solid"
        )

        total_fill = PatternFill(
            start_color="FFF2CC",
            end_color="FFF2CC",
            fill_type="solid"
        )

        center = Alignment(
            horizontal="center",
            vertical="center"
        )

        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        # ====================================
        # HEADER ROW
        # ====================================

        header_row = 3

        for col, header in enumerate(
            headers,
            start=1
        ):

            cell = ws.cell(
                row=header_row,
                column=col,
                value=header
            )

            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

        # ====================================
        # DATA ROWS
        # ====================================

        start_row = 4

        total_forecast = 0
        total_production = 0

        for row_idx, row in enumerate(
            inventory,
            start=start_row
        ):

            values = [
                row.get("product"),
                row.get("forecast"),
                round(row.get("bufferedDemand", 0), 2),
                row.get("band"),
                row.get("Z"),
                round(row.get("stdDev", 0), 2),
                round(row.get("safetyStock", 0), 2),
                round(row.get("minStock", 0), 2),
                round(row.get("maxStock", 0), 2),
                
            ]

            total_forecast += float(
                row.get("forecast", 0)
            )

            total_production += float(
                row.get("production", 0)
            )

            for col_idx, value in enumerate(
                values,
                start=1
            ):

                cell = ws.cell(
                    row=row_idx,
                    column=col_idx,
                    value=value
                )

                cell.font = data_font
                cell.alignment = center
                cell.border = border

        # ====================================
        # TOTAL ROW
        # ====================================

        total_row = start_row + len(inventory)

        totals = [
            "TOTAL",
            round(total_forecast, 2),
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            round(total_production, 2),
            ""
        ]

        for col_idx, value in enumerate(
            totals,
            start=1
        ):

            cell = ws.cell(
                row=total_row,
                column=col_idx,
                value=value
            )

            cell.font = total_font
            cell.fill = total_fill
            cell.alignment = center
            cell.border = border

        # ====================================
        # AUTO WIDTH
        # ====================================

        for col_num in range(
            1,
            ws.max_column + 1
        ):

            column_letter = get_column_letter(
                col_num
            )

            max_length = 0

            for row in range(
                1,
                ws.max_row + 1
            ):

                cell = ws.cell(
                    row=row,
                    column=col_num
                )

                try:

                    if cell.value:

                        max_length = max(
                            max_length,
                            len(str(cell.value))
                        )

                except:
                    pass

            ws.column_dimensions[
                column_letter
            ].width = max_length + 6

        # ====================================
        # ROW HEIGHT
        # ====================================

        for row in range(
            1,
            ws.max_row + 1
        ):

            ws.row_dimensions[
                row
            ].height = 28

        # ====================================
        # SAVE FILE
        # ====================================

        stream = BytesIO()

        wb.save(stream)

        stream.seek(0)

        return StreamingResponse(
            stream,

            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

            headers={
                "Content-Disposition":
                f"attachment; filename={selected_product}_inventory.xlsx"
            }
        )

    except Exception as e:

        print(
            "INVENTORY EXCEL ERROR:",
            str(e)
        )

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )

@app.post("/download-inventoryPivot-excel")
def download_inventory_excel(data: dict = Body(...)):

    from openpyxl import Workbook

    from openpyxl.styles import (
        Font,
        PatternFill,
        Alignment,
        Border,
        Side
    )

    from openpyxl.utils import get_column_letter

    from fastapi.responses import StreamingResponse

    from io import BytesIO

    try:

        inventory = data.get("inventory", [])

        selected_product = data.get(
            "selectedProduct",
            "inventory"
        )

        wb = Workbook()

        ws = wb.active

        ws.title = "Inventory Planning"

        # ====================================
        # TITLE
        # ====================================

        ws.merge_cells("A1:J1")

        title = ws["A1"]

        title.value = (
            f"{selected_product} Inventory Planning"
        )

        title.font = Font(
            bold=True,
            size=22
        )

        title.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        # ====================================
        # HEADERS
        # ====================================

        headers = [
            "Power (D)",
            "Sph",
            "Forecast",
            "Buffered",
            "Band",
            "Z",
            "Std Dev",
            "Safety",
            "Min",
            "Max"
        ]

        header_font = Font(
            bold=True,
            size=15
        )

        data_font = Font(
            size=13
        )

        total_font = Font(
            bold=True,
            size=14
        )

        header_fill = PatternFill(
            start_color="D9E1F2",
            end_color="D9E1F2",
            fill_type="solid"
        )

        total_fill = PatternFill(
            start_color="FFF2CC",
            end_color="FFF2CC",
            fill_type="solid"
        )

        center = Alignment(
            horizontal="center",
            vertical="center"
        )

        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        # ====================================
        # HEADER ROW
        # ====================================

        header_row = 3

        for col, header in enumerate(
            headers,
            start=1
        ):

            cell = ws.cell(
                row=header_row,
                column=col,
                value=header
            )

            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

        # ====================================
        # GROUP DATA
        # ====================================

        grouped = {}

        for row in inventory:

            key = str(
                row.get("sub_product", "")
            )

            if key not in grouped:
                grouped[key] = []

            grouped[key].append(row)

        # ====================================
        # WRITE DATA
        # ====================================

        current_row = 4

        total_forecast = 0

        for sub_product, items in grouped.items():

            start_merge_row = current_row

            for item in items:

                values = [
                    sub_product,
                    item.get("sub_sub_product"),
                    item.get("forecast"),
                    round(item.get("bufferedDemand", 0), 2),
                    item.get("band"),
                    item.get("Z"),
                    round(item.get("stdDev", 0), 2),
                    round(item.get("safetyStock", 0), 2),
                    round(item.get("minStock", 0), 2),
                    round(item.get("maxStock", 0), 2)
                ]

                total_forecast += float(
                    item.get("forecast", 0)
                )

                for col_idx, value in enumerate(
                    values,
                    start=1
                ):

                    cell = ws.cell(
                        row=current_row,
                        column=col_idx,
                        value=value
                    )

                    cell.font = data_font
                    cell.alignment = center
                    cell.border = border

                current_row += 1

            end_merge_row = current_row - 1

            # ====================================
            # MERGE POWER COLUMN
            # ====================================

            if start_merge_row != end_merge_row:

                ws.merge_cells(
                    start_row=start_merge_row,
                    start_column=1,
                    end_row=end_merge_row,
                    end_column=1
                )

                merged_cell = ws.cell(
                    row=start_merge_row,
                    column=1
                )

                merged_cell.alignment = center
                merged_cell.font = data_font
                merged_cell.border = border

        # ====================================
        # TOTAL ROW
        # ====================================

        total_row = current_row

        totals = [
            "TOTAL",
            "",
            round(total_forecast, 2),
            "",
            "",
            "",
            "",
            "",
            "",
            ""
        ]

        for col_idx, value in enumerate(
            totals,
            start=1
        ):

            cell = ws.cell(
                row=total_row,
                column=col_idx,
                value=value
            )

            cell.font = total_font
            cell.fill = total_fill
            cell.alignment = center
            cell.border = border

        # ====================================
        # AUTO WIDTH
        # ====================================

        for col_num in range(
            1,
            ws.max_column + 1
        ):

            column_letter = get_column_letter(
                col_num
            )

            max_length = 0

            for row in range(
                1,
                ws.max_row + 1
            ):

                cell = ws.cell(
                    row=row,
                    column=col_num
                )

                if cell.value:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            ws.column_dimensions[
                column_letter
            ].width = max_length + 6

        # ====================================
        # ROW HEIGHT
        # ====================================

        for row in range(
            1,
            ws.max_row + 1
        ):

            ws.row_dimensions[
                row
            ].height = 28

        # ====================================
        # SAVE
        # ====================================

        stream = BytesIO()

        wb.save(stream)

        stream.seek(0)

        return StreamingResponse(
            stream,

            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

            headers={
                "Content-Disposition":
                f"attachment; filename={selected_product}_inventory.xlsx"
            }
        )

    except Exception as e:

        print(
            "INVENTORY EXCEL ERROR:",
            str(e)
        )

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )

@app.post("/download-pivot-excel")
def download_pivot_excel(data: dict = Body(...)):

    from openpyxl import Workbook

    from openpyxl.styles import (
        Font,
        PatternFill,
        Alignment,
        Border,
        Side
    )

    from openpyxl.utils import (
        get_column_letter
    )

    from fastapi.responses import (
        StreamingResponse
    )

    from io import BytesIO

    try:

        results = data.get(
            "results",
            []
        )

        if not results:

            raise HTTPException(
                status_code=400,
                detail="No data"
            )

        # ====================================
        # ROWS / COLUMNS
        # ====================================

        rows = sorted(
    list(
        set(
            float(
                item.get(
                    "sub_product",
                    item.get("SUBP")
                )
            )
            for item in results
        )
    )
)

        columns = sorted(
    list(
        set(
            float(
                item.get(
                    "sub_sub_product",
                    item.get("SUBP2")
                )
            )
            for item in results
        )
    )
)

        # ====================================
        # WORKBOOK
        # ====================================

        wb = Workbook()

        ws = wb.active

        ws.title = "Pivot Forecast"

        # ====================================
        # TITLE
        # ====================================

        end_col = len(columns) + 2

        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=end_col
        )

        title = ws["A1"]

        title.value = "Forecast Allocation Pivot"

        title.font = Font(
            bold=True,
            size=20
        )

        title.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        # ====================================
        # STYLES
        # ====================================

        header_font = Font(
            bold=True,
            size=14
        )

        data_font = Font(
            size=12
        )

        total_font = Font(
            bold=True,
            size=13
        )

        header_fill = PatternFill(
            start_color="FFFF00",
            end_color="FFFF00",
            fill_type="solid"
        )

        total_fill = PatternFill(
            start_color="DCFCE7",
            end_color="DCFCE7",
            fill_type="solid"
        )

        center = Alignment(
            horizontal="center",
            vertical="center"
        )

        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        # ====================================
        # HEADER ROW
        # ====================================

        header_row = 3

        first = ws.cell(
            row=header_row,
            column=1,
            value="Power / Sph"
        )

        first.font = header_font
        first.fill = header_fill
        first.alignment = center
        first.border = border

        for idx, col_name in enumerate(
            columns,
            start=2
        ):

            cell = ws.cell(
                row=header_row,
                column=idx,
                value=col_name
            )

            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
            cell.border = border

        total_header = ws.cell(
            row=header_row,
            column=len(columns) + 2,
            value="Total"
        )

        total_header.font = header_font
        total_header.fill = header_fill
        total_header.alignment = center
        total_header.border = border

        # ====================================
        # DATA
        # ====================================

        current_row = 4

        grand_total = 0

        for row_name in rows:

            row_total = 0

            row_cell = ws.cell(
                row=current_row,
                column=1,
                value=str(row_name)
            )

            row_cell.font = total_font
            row_cell.alignment = center
            row_cell.border = border

            for col_idx, col_name in enumerate(
                columns,
                start=2
            ):

                entry = next(
    (
        r for r in results
        if float(
            r.get(
                "sub_product",
                r.get("SUBP")
            )
        ) == float(row_name)
        and
        float(
            r.get(
                "sub_sub_product",
                r.get("SUBP2")
            )
        ) == float(col_name)
    ),
    None
)

                value = (
    round(
        float(
            entry.get("forecast", 0)
        ),
        2
    )
    if entry
    else 0
)

                if value != "-":
                    row_total += float(value)
                    grand_total += float(value)
                cell = ws.cell(
                    row=current_row,
                    column=col_idx,
                    value=value
                )

                cell.font = data_font
                cell.alignment = center
                cell.border = border

            total_cell = ws.cell(
                row=current_row,
                column=len(columns) + 2,
                value=row_total
            )

            total_cell.font = total_font
            total_cell.fill = total_fill
            total_cell.alignment = center
            total_cell.border = border

            current_row += 1

        # ====================================
        # GRAND TOTAL
        # ====================================

        grand_row = current_row

        grand_label = ws.cell(
            row=grand_row,
            column=1,
            value="GRAND TOTAL"
        )

        grand_label.font = total_font
        grand_label.fill = total_fill
        grand_label.alignment = center
        grand_label.border = border

        grand_value = ws.cell(
            row=grand_row,
            column=len(columns) + 2,
            value=grand_total
        )

        grand_value.font = total_font
        grand_value.fill = total_fill
        grand_value.alignment = center
        grand_value.border = border

        # ====================================
        # AUTO WIDTH
        # ====================================

        for col_num in range(
            1,
            ws.max_column + 1
        ):

            letter = get_column_letter(
                col_num
            )

            max_length = 0

            for row in range(
                1,
                ws.max_row + 1
            ):

                cell = ws.cell(
                    row=row,
                    column=col_num
                )

                if cell.value:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            ws.column_dimensions[
                letter
            ].width = max_length + 5

        # ====================================
        # ROW HEIGHT
        # ====================================

        for row in range(
            1,
            ws.max_row + 1
        ):

            ws.row_dimensions[
                row
            ].height = 28

        # ====================================
        # SAVE
        # ====================================

        stream = BytesIO()

        wb.save(stream)

        stream.seek(0)

        return StreamingResponse(
            stream,

            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

            headers={
                "Content-Disposition":
                "attachment; filename=pivot_forecast.xlsx"
            }
        )

    except Exception as e:

        print(
            "PIVOT EXCEL ERROR:",
            str(e)
        )

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )
        
if __name__ == "__main__":

    import uvicorn

    uvicorn.run(
        app,
        host="0.0.0.0",
        port=5000
    )